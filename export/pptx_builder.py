"""
export/pptx_builder.py — assembles the .pptx from session state: KPIs,
chart images, insight text.

IMPORTANT: this module must never import insight.client (the OpenRouter
wrapper). It only reads pre-generated insight text from insight.cache. This
is enforced by convention here and checked in
tests/test_export_no_network.py by asserting "insight.client" never appears
in this module's import graph.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import datetime, timezone

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from export.chart_images import figure_to_png_stream
from insight import cache as insight_cache

SLIDE_WIDTH_IN = 13.333
SLIDE_HEIGHT_IN = 7.5

COLOR_BG = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_TEXT = RGBColor(0x1A, 0x1A, 0x1A)
COLOR_MUTED = RGBColor(0x66, 0x66, 0x66)
COLOR_ACCENT_POS = RGBColor(0x1E, 0x7A, 0x34)
COLOR_ACCENT_NEG = RGBColor(0xB0, 0x2A, 0x2A)
COLOR_ACCENT = RGBColor(0x2A, 0x4B, 0xB0)

FOOTNOTE_TEXT = (
    "Generated locally from the uploaded data. AI-assisted narrative insights via "
    "OpenRouter; underlying figures computed locally."
)


@dataclass
class KPITileData:
    label: str
    value_display: str
    delta_display: str | None = None
    delta_positive: bool | None = None


@dataclass
class ViewExportData:
    """Everything one dashboard view (e.g. Executive, Revenue-equivalent)
    needs to render its export slide(s). Built by ui/*.py from session state
    — pptx_builder never recomputes analytics itself."""

    view_name: str
    kpi_tiles: list[KPITileData] = field(default_factory=list)
    trend_figure: object | None = None  # a Plotly Figure, or None
    drivers_table: list[dict] | None = None  # [{category, contribution_pct, ...}]
    drivers_figure: object | None = None


@dataclass
class DeckInputs:
    dataset_name: str
    period_covered: str
    views: list[ViewExportData]
    risk_signals: list[dict]  # [{description, severity}]


def _blank_slide(prs: Presentation):
    blank_layout = prs.slide_layouts[6]
    return prs.slides.add_slide(blank_layout)


def _add_title_text(slide, text: str, top: float = 0.4, size: int = 32, bold: bool = True, color=COLOR_TEXT):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(SLIDE_WIDTH_IN - 1.0), Inches(1.0))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def _add_body_text(slide, text: str, left: float, top: float, width: float, height: float, size: int = 14, color=COLOR_TEXT, bold: bool = False):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def _add_kpi_tile(slide, tile: KPITileData, left: float, top: float, width: float = 2.6, height: float = 1.5):
    box = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))  # 1 = RECTANGLE
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0xF4, 0xF6, 0xFB)
    box.line.color.rgb = RGBColor(0xDD, 0xE2, 0xEE)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(91440)
    tf.margin_top = Emu(45720)

    p_label = tf.paragraphs[0]
    r_label = p_label.add_run()
    r_label.text = tile.label
    r_label.font.size = Pt(11)
    r_label.font.color.rgb = COLOR_MUTED

    p_value = tf.add_paragraph()
    r_value = p_value.add_run()
    r_value.text = tile.value_display
    r_value.font.size = Pt(22)
    r_value.font.bold = True
    r_value.font.color.rgb = COLOR_TEXT

    if tile.delta_display:
        p_delta = tf.add_paragraph()
        r_delta = p_delta.add_run()
        r_delta.text = tile.delta_display
        r_delta.font.size = Pt(12)
        r_delta.font.bold = True
        r_delta.font.color.rgb = (
            COLOR_ACCENT_POS if tile.delta_positive else COLOR_ACCENT_NEG
        ) if tile.delta_positive is not None else COLOR_MUTED

    return box


def _add_picture_from_figure(slide, fig, left: float, top: float, width: float):
    stream = figure_to_png_stream(fig)
    slide.shapes.add_picture(stream, Inches(left), Inches(top), width=Inches(width))


def _cached_insight_for_view(view_name: str) -> dict | None:
    return insight_cache.get_for_view(view_name)


def build_title_slide(prs: Presentation, dataset_name: str, period_covered: str) -> None:
    slide = _blank_slide(prs)
    _add_title_text(slide, "Aivara Insight Lite", top=2.4, size=40)
    _add_body_text(
        slide,
        f"Dataset: {dataset_name}",
        left=0.5,
        top=3.3,
        width=SLIDE_WIDTH_IN - 1.0,
        height=0.5,
        size=16,
        color=COLOR_MUTED,
    )
    _add_body_text(
        slide,
        f"Period covered: {period_covered}",
        left=0.5,
        top=3.8,
        width=SLIDE_WIDTH_IN - 1.0,
        height=0.5,
        size=16,
        color=COLOR_MUTED,
    )
    generated_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    _add_body_text(
        slide,
        f"Generated {generated_at}",
        left=0.5,
        top=4.3,
        width=SLIDE_WIDTH_IN - 1.0,
        height=0.5,
        size=12,
        color=COLOR_MUTED,
    )


def build_executive_summary_slide(prs: Presentation, views: list[ViewExportData]) -> None:
    slide = _blank_slide(prs)
    _add_title_text(slide, "Executive Summary")

    left = 0.5
    top = 1.6
    tile_width = 2.9
    gap = 0.25
    all_tiles = [t for v in views for t in v.kpi_tiles][:4]
    for i, tile in enumerate(all_tiles):
        _add_kpi_tile(slide, tile, left + i * (tile_width + gap), top, width=tile_width)

    if views:
        primary_view = views[0]
        insight = _cached_insight_for_view(primary_view.view_name)
        if insight and insight.get("headline"):
            _add_body_text(
                slide,
                insight["headline"],
                left=0.5,
                top=3.6,
                width=SLIDE_WIDTH_IN - 1.0,
                height=1.2,
                size=18,
                bold=True,
                color=COLOR_ACCENT,
            )


def build_view_slide(prs: Presentation, view: ViewExportData) -> None:
    slide = _blank_slide(prs)
    _add_title_text(slide, view.view_name)

    left = 0.5
    top = 1.3
    tile_width = 2.3
    gap = 0.2
    for i, tile in enumerate(view.kpi_tiles[:4]):
        _add_kpi_tile(slide, tile, left + i * (tile_width + gap), top, width=tile_width, height=1.2)

    chart_top = 2.8 if view.kpi_tiles else 1.4
    insight = _cached_insight_for_view(view.view_name)
    chart_width = 7.5 if insight else 11.5
    if view.trend_figure is not None:
        try:
            _add_picture_from_figure(slide, view.trend_figure, left=0.5, top=chart_top, width=chart_width)
        except Exception:
            pass  # fail soft — omit the picture rather than crash export
    if insight:
        text_left = 8.3
        text_width = SLIDE_WIDTH_IN - text_left - 0.5
        y = chart_top
        if insight.get("headline"):
            _add_body_text(slide, insight["headline"], text_left, y, text_width, 0.9, size=14, bold=True)
            y += 0.9
        if insight.get("driver_explanation"):
            _add_body_text(slide, insight["driver_explanation"], text_left, y, text_width, 0.9, size=12)
            y += 0.9
        if insight.get("suggested_action"):
            _add_body_text(
                slide, "Next: " + insight["suggested_action"], text_left, y, text_width, 0.9, size=12, color=COLOR_ACCENT
            )
    # If insight wasn't generated for this view, we simply omit the text
    # block rather than showing a placeholder (Section 8.7).


def build_drivers_slide(prs: Presentation, view: ViewExportData) -> None:
    if view.drivers_figure is None and not view.drivers_table:
        return
    slide = _blank_slide(prs)
    _add_title_text(slide, f"Drivers — {view.view_name}")

    if view.drivers_figure is not None:
        try:
            _add_picture_from_figure(slide, view.drivers_figure, left=0.5, top=1.4, width=7.0)
        except Exception:
            pass

    if view.drivers_table:
        text_left = 7.8
        y = 1.4
        _add_body_text(slide, "Top contributors", text_left, y, 5.0, 0.5, size=14, bold=True)
        y += 0.6
        for row in view.drivers_table[:8]:
            category = row.get("category", "")
            pct = row.get("contribution_pct", 0)
            color = COLOR_ACCENT_POS if pct >= 0 else COLOR_ACCENT_NEG
            _add_body_text(slide, f"{category}: {pct:+.1f}%", text_left, y, 5.0, 0.4, size=12, color=color)
            y += 0.4


def build_risk_signals_slide(prs: Presentation, risk_signals: list[dict]) -> None:
    if not risk_signals:
        return  # never ship an empty slide
    slide = _blank_slide(prs)
    _add_title_text(slide, "Risk Signals")

    y = 1.5
    for signal in risk_signals[:10]:
        severity = signal.get("severity", "low")
        color = {"high": COLOR_ACCENT_NEG, "medium": RGBColor(0xB8, 0x86, 0x00)}.get(severity, COLOR_MUTED)
        _add_body_text(
            slide,
            f"[{severity.upper()}] {signal.get('description', '')}",
            left=0.5,
            top=y,
            width=SLIDE_WIDTH_IN - 1.0,
            height=0.5,
            size=13,
            color=color,
        )
        y += 0.55


def build_closing_slide(prs: Presentation) -> None:
    slide = _blank_slide(prs)
    _add_title_text(slide, "About this report", top=2.6, size=24)
    _add_body_text(
        slide,
        FOOTNOTE_TEXT,
        left=0.5,
        top=3.4,
        width=SLIDE_WIDTH_IN - 1.0,
        height=1.2,
        size=13,
        color=COLOR_MUTED,
    )


def build_deck(inputs: DeckInputs) -> io.BytesIO:
    """Builds the full deck in-memory (Section 8.7): title, executive
    summary, one slide per view, drivers, risk (only if any), closing. Never
    writes to disk — returns a BytesIO the caller passes straight to
    st.download_button. Never calls the OpenRouter client — only reads from
    insight.cache."""
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_WIDTH_IN)
    prs.slide_height = Inches(SLIDE_HEIGHT_IN)

    build_title_slide(prs, inputs.dataset_name, inputs.period_covered)
    build_executive_summary_slide(prs, inputs.views)

    for view in inputs.views:
        build_view_slide(prs, view)
        build_drivers_slide(prs, view)

    build_risk_signals_slide(prs, inputs.risk_signals)
    build_closing_slide(prs)

    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer
